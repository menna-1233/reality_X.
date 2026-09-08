"""Export reports to an Excel (.xlsx) file for offline review or handoff.

Used by GET /reports/export/excel — builds a right-to-left spreadsheet
with Arabic headers and one row per report.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADERS = [
    "رقم البلاغ",
    "تاريخ البلاغ",
    "نوع المشكلة",
    "درجة الخطورة",
    "الإدارة المسؤولة",
    "الوصف",
    "الموقع",
    "ملخص AI",
    "نسبة الثقة",
    "الحالة",
    "رابط الصورة",
]

_PROBLEM_TYPE_AR = {
    "pothole": "حفرة",
    "garbage": "قمامة",
    "water_leak": "تسريب مياه",
    "broken_light": "إنارة معطلة",
    "accident": "حادث",
    "other": "أخرى",
}

_SEVERITY_AR = {"low": "منخفضة", "medium": "متوسطة", "high": "عالية", "critical": "حرجة"}

_STATUS_AR = {
    "open": "مفتوح",
    "in_progress": "قيد المعالجة",
    "resolved": "تم الحل",
    "closed": "مغلق",
}

_COLUMN_WIDTHS = [12, 18, 14, 12, 24, 35, 20, 35, 10, 14, 40]


def build_reports_excel(reports: list[dict]) -> bytes:
    """Build an .xlsx workbook from a list of report dicts. Returns raw bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "البلاغات"
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col_idx, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, report in enumerate(reports, start=2):
        confidence = report.get("confidence")
        values = [
            (report.get("id") or "")[:8],
            (report.get("created_at") or "")[:19].replace("T", " "),
            _PROBLEM_TYPE_AR.get(report.get("problem_type"), report.get("problem_type") or ""),
            _SEVERITY_AR.get(report.get("severity"), report.get("severity") or ""),
            report.get("department") or "",
            report.get("description") or "",
            report.get("location_text") or "",
            report.get("ai_summary") or "",
            f"{confidence * 100:.0f}%" if confidence is not None else "",
            _STATUS_AR.get(report.get("status"), report.get("status") or ""),
            report.get("image_url") or "",
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)

    for col_idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
